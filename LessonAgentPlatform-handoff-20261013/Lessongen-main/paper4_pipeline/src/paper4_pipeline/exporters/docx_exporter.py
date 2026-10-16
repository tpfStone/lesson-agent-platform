"""Print-friendly A4 Word export for a structured lesson plan."""

from __future__ import annotations

import json
import math
import os
import tempfile
from pathlib import Path

from paper4_pipeline.domain.models import LessonPlanDocument, PipelineResult
from paper4_pipeline.exporters.json_exporter import best_version


TABLE_WIDTH_DXA = 9500
BLACK = "202020"
DARK_GRAY = "555555"
MID_GRAY = "D9D9D9"
LIGHT_GRAY = "F2F2F2"


def export_best_docx(result: PipelineResult, path: Path) -> Path:
    version = best_version(result)
    return export_docx(
        version.document,
        path,
        run_id=result.run_id,
        best_version_id=result.best_version_id,
        last_version_id=result.last_version_id,
    )


def export_docx(
    document: LessonPlanDocument,
    path: Path,
    *,
    run_id: str,
    best_version_id: str,
    last_version_id: str,
) -> Path:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Word export requires python-docx; install the optional 'docx' extra."
        ) from exc

    def set_family(r_pr, family: str) -> None:
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        for attribute in list(r_fonts.attrib):
            if attribute.lower().endswith("theme"):
                del r_fonts.attrib[attribute]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn("w:" + script), family)

    def set_font(
        run,
        *,
        size: float = 10.5,
        bold: bool | None = None,
        color: str = BLACK,
        chinese: str = "SimSun",
    ) -> None:
        set_family(run._element.get_or_add_rPr(), chinese)
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    def set_cell_shading(cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def set_cell_width(cell, width: int) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(width))
        tc_w.set(qn("w:type"), "dxa")

    def set_cell_margins(cell, horizontal: int = 110, vertical: int = 75) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for edge, value in (
            ("top", vertical),
            ("bottom", vertical),
            ("left", horizontal),
            ("right", horizontal),
        ):
            node = tc_mar.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_repeat_header(row) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

    def configure_table(table, widths: list[int]) -> None:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(sum(widths)))
        tbl_w.set(qn("w:type"), "dxa")
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(width))
            grid.append(col)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                set_cell_width(cell, widths[min(index, len(widths) - 1)])
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    def paragraph_in(cell, *, align=None):
        paragraph = cell.add_paragraph() if cell.text else cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.widow_control = True
        paragraph.alignment = align or WD_ALIGN_PARAGRAPH.LEFT
        return paragraph

    def clear_cell(cell) -> None:
        cell.text = ""

    def add_cell_text(
        cell,
        text: str,
        *,
        label: str = "",
        bold: bool = False,
        size: float = 9.5,
        align=None,
        chinese: str = "SimSun",
    ) -> None:
        paragraph = paragraph_in(cell, align=align)
        if label:
            label_run = paragraph.add_run(label)
            set_font(label_run, size=size, bold=True)
        run = paragraph.add_run(text or "—")
        set_font(run, size=size, bold=bold, chinese=chinese)

    def add_cell_list(cell, label: str, values: list[str], *, size: float = 9.5) -> None:
        if not values:
            return
        heading = paragraph_in(cell)
        set_font(heading.add_run(label), size=size, bold=True)
        for value in values:
            p = paragraph_in(cell)
            p.paragraph_format.left_indent = Pt(8)
            p.paragraph_format.first_line_indent = Pt(-8)
            set_font(p.add_run(f"• {value}"), size=size)

    section_number = 0

    def add_section_bar(title: str) -> None:
        nonlocal section_number
        section_number += 1
        paragraph = doc.add_paragraph(style="Heading 1")
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        set_font(
            paragraph.add_run(f"{_cn_number(section_number)}、{title}"),
            size=11,
            bold=True,
            chinese="SimHei",
        )

    def add_body(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.widow_control = True
        set_font(p.add_run(text), size=10.5)

    def add_bullets(values: list[str]) -> None:
        for value in values:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-12)
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(f"• {value}"), size=10.5)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    set_family(normal._element.get_or_add_rPr(), "SimSun")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.2

    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        doc.styles.element.insert(0, defaults)
    r_default = defaults.find(qn("w:rPrDefault"))
    if r_default is None:
        r_default = OxmlElement("w:rPrDefault")
        defaults.append(r_default)
    r_pr = r_default.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        r_default.append(r_pr)
    set_family(r_pr, "SimSun")
    for name in ("Title", *(f"Heading {index}" for index in range(1, 10))):
        style = doc.styles[name]
        set_family(style.element.get_or_add_rPr(), "SimHei")
        style.font.color.rgb = RGBColor.from_string(BLACK)
        border = style.element.find("w:pPr/w:pBdr", style.element.nsmap)
        if border is not None:
            border.getparent().remove(border)

    doc.core_properties.title = f"{document.metadata.topic}教学设计"
    doc.core_properties.subject = "结构化教学设计"
    doc.core_properties.author = "Paper#4 多智能体教案流水线"
    doc.core_properties.comments = "AI 生成草案，正式使用前须由教师核验。"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        header.add_run(
            f"教学设计 · {document.metadata.subject} / {document.metadata.grade}"
        ),
        size=8,
        color=DARK_GRAY,
        chinese="SimSun",
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("AI 生成草案 · 须教师审核  |  "), size=8, color=DARK_GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    set_font(
        title.add_run(f"{document.metadata.topic}教学设计"),
        size=18,
        bold=True,
        chinese="SimHei",
    )

    metadata = doc.add_table(rows=2, cols=4)
    configure_table(metadata, [2375, 2375, 2375, 2375])
    labels = ["学科", "年级", "课题", "课时"]
    values = [
        document.metadata.subject,
        document.metadata.grade,
        document.metadata.topic,
        f"{document.metadata.duration_minutes} 分钟",
    ]
    for index, label in enumerate(labels):
        clear_cell(metadata.cell(0, index))
        set_cell_shading(metadata.cell(0, index), LIGHT_GRAY)
        add_cell_text(
            metadata.cell(0, index),
            label,
            bold=True,
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            chinese="SimSun",
        )
        clear_cell(metadata.cell(1, index))
        add_cell_text(
            metadata.cell(1, index),
            values[index],
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    if not document.curriculum_standards:
        add_body("未提供课程标准原文，使用前须由教师核验。")
    sections: list[tuple[str, str | list[str]]] = [
        ("课程标准依据", document.curriculum_standards),
        ("教学内容分析", document.content_analysis),
        ("学情分析", document.student_analysis),
    ]
    for heading, value in sections:
        if not value:
            continue
        add_section_bar(heading)
        add_bullets(value) if isinstance(value, list) else add_body(value)

    if document.design_thesis or document.driving_question or document.learning_trajectory:
        add_section_bar("教学设计主线")
        if document.design_thesis:
            add_body("设计主张：" + document.design_thesis)
        if document.driving_question:
            add_body("驱动问题：" + document.driving_question)
        if document.learning_trajectory:
            add_body("学习轨迹：" + " → ".join(document.learning_trajectory))

    if document.learning_objectives:
        add_section_bar("学习目标与达成证据")
        table = doc.add_table(rows=1, cols=3)
        configure_table(table, [1000, 4100, 4400])
        for idx, label in enumerate(["编号", "学习目标", "达成证据"]):
            clear_cell(table.cell(0, idx))
            set_cell_shading(table.cell(0, idx), LIGHT_GRAY)
            add_cell_text(table.cell(0, idx), label, bold=True)
        set_repeat_header(table.rows[0])
        for objective in document.learning_objectives:
            row = table.add_row()
            for idx, width in enumerate([1000, 4100, 4400]):
                set_cell_width(row.cells[idx], width)
                set_cell_margins(row.cells[idx])
                clear_cell(row.cells[idx])
            add_cell_text(row.cells[0], objective.objective_id)
            add_cell_text(row.cells[1], objective.description)
            add_cell_text(row.cells[2], objective.evidence_of_achievement or "待补充")

    if document.key_points or document.difficult_points or document.teaching_strategy:
        add_section_bar("教学重点、难点与策略")
    if document.key_points:
        add_body("教学重点：" + "；".join(document.key_points))
    if document.difficult_points:
        add_body("教学难点：" + "；".join(document.difficult_points))
    if document.teaching_strategy:
        add_body("教学策略：" + document.teaching_strategy)

    if document.resources:
        add_section_bar("教学准备")
        for resource in document.resources:
            add_body(resource.name + (f"：{resource.description}" if resource.description else ""))
            if resource.ready_to_use_content:
                add_body("材料内容：" + resource.ready_to_use_content)

    add_section_bar("教学过程")
    resource_names = {item.resource_id: item.name for item in document.resources}
    artifact_titles = {
        item.artifact_id: item.title for item in document.teaching_artifacts
    }
    process = doc.add_table(rows=1, cols=3)
    process_widths = [3900, 3300, 2300]
    configure_table(process, process_widths)
    for idx, label in enumerate(
        ["教学环节与教师活动", "学生学习任务与产出", "评价重点与调控"]
    ):
        clear_cell(process.cell(0, idx))
        set_cell_shading(process.cell(0, idx), LIGHT_GRAY)
        add_cell_text(
            process.cell(0, idx), label, bold=True, size=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
    set_repeat_header(process.rows[0])

    for index, step in enumerate(document.procedure_steps, start=1):
        row = process.add_row()
        for col, width in enumerate(process_widths):
            set_cell_width(row.cells[col], width)
            set_cell_margins(row.cells[col], horizontal=90, vertical=65)
            row.cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            clear_cell(row.cells[col])

        add_cell_text(
            row.cells[0], f"{index}. {step.stage}（{step.duration_minutes}分钟）",
            bold=True
        )
        add_cell_list(row.cells[0], "教师活动", step.teacher_actions)
        for q_index, question in enumerate(step.questions, start=1):
            add_cell_text(row.cells[0], question.question, label=f"问题{q_index}：")
            add_cell_list(row.cells[0], "预期回答", question.expected_responses)
            add_cell_list(row.cells[0], "追问", question.teacher_follow_ups)
        if step.transition:
            add_cell_text(row.cells[0], step.transition, label="环节过渡：")
        if step.resource_ids:
            add_cell_list(
                row.cells[0],
                "使用资源",
                [resource_names[item] for item in step.resource_ids],
            )
        if step.artifact_ids:
            add_cell_list(
                row.cells[0],
                "使用材料",
                [artifact_titles[item] for item in step.artifact_ids],
            )

        add_cell_list(row.cells[1], "学生任务", step.student_actions)
        if step.student_product:
            add_cell_text(row.cells[1], step.student_product, label="学习产出：")
        add_cell_list(row.cells[1], "成功标准", step.success_criteria)
        add_cell_list(row.cells[1], "基础支架", step.scaffolds)
        add_cell_list(row.cells[1], "拓展挑战", step.extensions)

        if step.assessment:
            add_cell_text(row.cells[2], step.assessment, label="评价证据：")
        misconceptions = [
            item for question in step.questions for item in question.possible_misconceptions
        ]
        evidence = [q.evidence_to_notice for q in step.questions if q.evidence_to_notice]
        add_cell_list(row.cells[2], "观察重点", evidence)
        add_cell_list(row.cells[2], "误概念预判", misconceptions)
        for branch in step.response_branches:
            add_cell_text(
                row.cells[2],
                f"若{branch.trigger}，则{branch.teacher_move}"
                + (f"（{branch.purpose}）" if branch.purpose else ""),
                label="调控：",
            )
        if step.design_rationale:
            add_cell_text(row.cells[2], step.design_rationale, label="设计意图：")

    if document.teaching_artifacts:
        add_section_bar("可直接使用的教学材料")
        for index, artifact in enumerate(document.teaching_artifacts, start=1):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            set_font(
                p.add_run(f"材料 {index}｜{artifact.title}（{artifact.artifact_type}）"),
                size=11, bold=True, chinese="SimHei"
            )
            material = doc.add_table(rows=0, cols=2)
            configure_table(material, [1550, 7950])
            for label, value in (
                ("用途", artifact.purpose),
                ("材料正文", artifact.content),
                ("答案/标准", artifact.answer_or_success_criteria),
            ):
                if not value:
                    continue
                row = material.add_row()
                for idx, width in enumerate([1550, 7950]):
                    set_cell_width(row.cells[idx], width)
                    set_cell_margins(row.cells[idx])
                    clear_cell(row.cells[idx])
                set_cell_shading(row.cells[0], LIGHT_GRAY)
                add_cell_text(row.cells[0], label, bold=True)
                add_cell_text(row.cells[1], value)

    for title_text, value in (
        ("总体评价设计", document.assessment_plan),
        ("差异化支持", document.differentiation),
        ("作业设计", document.homework),
        ("板书设计", document.board_design),
        ("课后观察与反思", document.reflection),
    ):
        if value.strip():
            add_section_bar(title_text)
            add_body(value)
    if document.references:
        add_section_bar("参考材料")
        add_bullets(document.references)

    if document.template_extensions:
        add_section_bar("补充内容")
        for key, value in document.template_extensions.items():
            if key == "source_table_columns" and isinstance(value, list):
                for column in value:
                    add_body(f"原稿补充列：{column['heading']}（{column['table_locator']}，第 {column['column_index'] + 1} 列）")
                    notes = doc.add_table(rows=1, cols=2)
                    configure_table(notes, [2700, 6800])
                    for index, label in enumerate(["教学环节", column["heading"]]):
                        add_cell_text(notes.cell(0, index), label, bold=True)
                        set_cell_shading(notes.cell(0, index), LIGHT_GRAY)
                    set_repeat_header(notes.rows[0])
                    for source_row in column["rows"]:
                        row = notes.add_row()
                        add_cell_text(row.cells[0], source_row["stage"])
                        add_cell_text(row.cells[1], source_row["text"])
                continue
            add_body({"extra_notes": "额外备注"}.get(key, key) + "：")
            if isinstance(value, str):
                add_body(value)
            elif isinstance(value, list) and value and all(isinstance(item, str) for item in value):
                add_bullets(value)
            else:
                # Unknown extension schemas remain visible without flattening
                # away their keys, ordering, or false/zero/null values.
                add_body(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False))

    add_section_bar("教师核验与实施记录")
    review = doc.add_table(rows=0, cols=2)
    configure_table(review, [2700, 6800])
    for label, value in (
        ("使用前核验", "□ 课标与教材原文  □ 例题/语料与答案  □ 时间与资源  □ 安全、伦理与隐私"),
        ("实施证据", "记录学生关键回答、作品差异、误概念以及教师临场调整。"),
        ("下次修订", "保留：　　　　　　　　　调整：　　　　　　　　　删除/新增："),
    ):
        row = review.add_row()
        for idx, width in enumerate([2700, 6800]):
            set_cell_width(row.cells[idx], width)
            set_cell_margins(row.cells[idx], vertical=100)
            clear_cell(row.cells[idx])
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        add_cell_text(row.cells[0], label, bold=True)
        add_cell_text(row.cells[1], value)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_together = True
    set_font(
        p.add_run(
            f"生成记录：run={run_id}；best={best_version_id}；last={last_version_id}。"
            "本教案为 AI 生成草案，课程标准、教材内容与答案须由教师核验。"
        ),
        size=8, color=DARK_GRAY, chinese="SimSun"
    )

    # Keep compact rows readable across page boundaries. Long cells must remain
    # splittable; a blanket cantSplit makes a page-sized row overflow or leave gaps.
    for table in doc.tables:
        for row in table.rows:
            estimated_lines = []
            for cell in row.cells:
                width = int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
                characters_per_line = max(1, width / 20 / 9.5 - 2)
                estimated_lines.append(sum(max(1, math.ceil(len(line) / characters_per_line))
                                           for line in cell.text.split("\n")))
            if max(estimated_lines, default=0) <= 14:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{path.stem}.", suffix=".docx", dir=path.parent
    )
    os.close(descriptor)
    try:
        doc.save(temporary)
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except OSError:
            pass
        raise
    return path


def _cn_number(value: int) -> str:
    mapping = {
        1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六",
        7: "七", 8: "八", 9: "九", 10: "十", 11: "十一", 12: "十二",
        13: "十三", 14: "十四", 15: "十五",
    }
    return mapping.get(value, str(value))
