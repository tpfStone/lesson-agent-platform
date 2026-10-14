"""Offline boundary tests; no provider, credentials, or artifact registry required."""
import hashlib
from pathlib import Path

from docx import Document
from paper4_pipeline.web_api.docx_ingestion import extract_docx
from paper4_pipeline.web_api.generated_docx import restore_generated_docx
from paper4_pipeline.web_api.schemas import EngineLessonInput
from paper4_pipeline.web_api.settings import EngineSettings


def settings(tmp_path):
    return EngineSettings(state_root=tmp_path / 'state', artifacts_root=tmp_path / 'artifacts',
                          config_path=tmp_path / 'unused-config.json')


def make_document(tmp_path, teacher='引导比较'):
    doc = Document()
    doc.add_paragraph('教学目标：理解面积')
    table = doc.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ['环节', '时间', '教师活动', '学生活动', '评价']):
        cell.text = value
    for index, minutes in enumerate([5, 12, 8, 9, 4, 2]):
        for cell, value in zip(table.add_row().cells, [f'环节{index + 1}', str(minutes), teacher,
                                                     '学生探究\n表达结论', '观察并反馈']):
            cell.text = value
    doc.add_paragraph('作业：测量三角形')
    path = tmp_path / 'current-version.docx'
    doc.save(path)
    return path


def test_current_docx_retains_six_rows_full_cells_and_nonclassroom_content(tmp_path):
    path = make_document(tmp_path)
    original = path.read_bytes()
    raw = extract_docx(path, original_filename=path.name, settings=settings(tmp_path))
    assert raw.source_sha256 == hashlib.sha256(original).hexdigest()
    assert path.read_bytes() == original
    assert raw.tables[0].rows == 7 and raw.tables[0].columns == 5
    assert [cell.text for cell in raw.tables[0].cells if cell.column == 1 and cell.row > 0] == ['5', '12', '8', '9', '4', '2']
    assert raw.ordered_blocks[0].text == '教学目标：理解面积'
    assert raw.ordered_blocks[-1].text == '作业：测量三角形'
    assert sum('学生探究\n表达结论' in block.text for block in raw.ordered_blocks) == 6


def test_edited_docx_does_not_claim_a_trusted_generated_artifact(tmp_path):
    path = make_document(tmp_path, '教师确认的修改')
    config = settings(tmp_path)
    raw = extract_docx(path, original_filename=path.name, settings=config)
    task = EngineLessonInput(mode='optimize', subject='数学', grade='五年级', topic='面积')
    assert restore_generated_docx(raw, task, task_id='task', plan_id='plan', artifacts_root=config.artifacts_root) is None
    assert any(cell.text == '教师确认的修改' for cell in raw.tables[0].cells)
