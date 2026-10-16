"""The Word template must not create empty sections or forbid long cell pagination."""
from docx import Document

from paper4_pipeline.domain.models import LessonPlanDocument
from paper4_pipeline.exporters.docx_exporter import export_docx


def export(tmp_path, **values):
    teacher = values.pop("teacher_text", "比较面积")
    lesson = LessonPlanDocument(
        plan_id="layout", task_id="layout",
        metadata={"subject": "数学", "grade": "五年级", "topic": "面积", "duration_minutes": 40},
        procedure_steps=[{"step_id": "s1", "stage": "探究", "duration_minutes": 40,
                          "teacher_actions": [teacher], "student_actions": ["说明理由"],
                          "assessment": "解释是否合理"}],
        **values,
    )
    path = export_docx(lesson, tmp_path / "layout.docx", run_id="layout", best_version_id="v0", last_version_id="v0")
    return Document(path)


def test_absent_content_has_no_empty_or_invented_sections(tmp_path):
    doc = export(tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    # Optional fields in a local revision are not a license to invent a full lesson.
    assert "教学重点、难点与策略" not in text
    assert "教学设计主线" not in text
    assert "教学内容分析" not in text
    assert "学情分析" not in text
    assert "课程标准依据" not in text
    assert "未提供课程标准原文" in text
    assert "一、教学过程" in text
    assert "二、教师核验与实施记录" in text


def test_long_objective_and_material_cells_can_flow_across_pages(tmp_path):
    doc = export(tmp_path,
        learning_objectives=[{"objective_id": "obj", "description": "长目标。" * 800,
                              "evidence_of_achievement": "证据"}],
        teaching_artifacts=[{"artifact_id": "a1", "artifact_type": "worksheet",
                             "title": "工作单", "purpose": "探究", "content": "长材料。" * 800,
                             "answer_or_success_criteria": "合理解释"}],
    )
    long_rows = [row for table in doc.tables for row in table.rows
                 if any(len(cell.text) > 1000 for cell in row.cells)]
    assert len(long_rows) == 2
    assert all(row._tr.find("w:trPr/w:cantSplit", row._tr.nsmap) is None for row in long_rows)
    assert "长目标。" * 800 in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_section_headings_are_unboxed_and_numbering_is_contiguous(tmp_path):
    doc = export(tmp_path, content_analysis="比较面积单位", homework="说明计算方法")
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert [p.text for p in headings] == ["一、教学内容分析", "二、教学过程", "三、作业设计", "四、教师核验与实施记录"]
    assert all(p._p.find("w:pPr/w:pBdr", p._p.nsmap) is None for p in headings)
    assert all(p.paragraph_format.keep_with_next for p in headings)
    title_style = doc.styles["Title"]
    assert title_style.element.find("w:pPr/w:pBdr", title_style.element.nsmap) is None


def test_short_teaching_rows_stay_together_but_long_rows_can_continue(tmp_path):
    doc = export(tmp_path)
    process = next(t for t in doc.tables if "教师活动" in t.cell(0, 0).text)
    assert process.rows[1]._tr.find("w:trPr/w:cantSplit", process.rows[1]._tr.nsmap) is not None
    long_doc = export(tmp_path, teacher_text="完整的长教师活动。" * 800)
    long_process = next(t for t in long_doc.tables if "教师活动" in t.cell(0, 0).text)
    assert long_process.rows[1]._tr.find("w:trPr/w:cantSplit", long_process.rows[1]._tr.nsmap) is None


def test_fonts_cover_all_scripts_and_editable_defaults_without_changing_table_roles(tmp_path):
    from docx.oxml.ns import qn
    doc = export(tmp_path, content_analysis="面积 Area 12 cm²", reflection="备注 Notes 123",
                 template_extensions={"extra_notes": ["保留中文与🙂"]})
    def check(r_pr, expected):
        assert r_pr is not None
        fonts = r_pr.find(qn("w:rFonts"))
        assert fonts is not None
        assert {fonts.get(qn("w:" + slot)) for slot in ("ascii", "hAnsi", "eastAsia", "cs")} == {expected}
        assert not any(key.lower().endswith("theme") for key in fonts.attrib)
    for paragraph in doc.paragraphs:
        expected = "SimHei" if paragraph.style.name in {"Title", "Heading 1"} else "SimSun"
        for run in paragraph.runs:
            if run.text:
                check(run._element.rPr, expected)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text:
                            check(run._element.rPr, "SimSun")
    for section in doc.sections:
        for region in (section.header, section.footer):
            for paragraph in region.paragraphs:
                for run in paragraph.runs:
                    if run.text:
                        check(run._element.rPr, "SimSun")
    for name, family in [("Normal", "SimSun"), ("Title", "SimHei"), ("Heading 1", "SimHei")]:
        check(doc.styles[name].element.rPr, family)
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    check(defaults, "SimSun")
    process = next(t for t in doc.tables if "教师活动" in t.cell(0, 0).text)
    assert all(run.bold for cell in process.rows[0].cells for run in cell.paragraphs[0].runs if run.text)
