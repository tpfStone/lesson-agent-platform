"""Exported Word files must retain canonical extension content without model calls."""
from datetime import datetime, timezone
import hashlib
import json

from docx import Document
import pytest

from paper4_pipeline.domain.models import LessonPlanDocument, LessonPlanVersion, PipelineResult
from paper4_pipeline.exporters.docx_exporter import export_docx
from paper4_pipeline.exporters.manifest import export_artifacts


@pytest.fixture
def lesson():
    return LessonPlanDocument(
        plan_id="plan", task_id="task",
        metadata={"subject": "数学", "grade": "三年级", "topic": "面积", "duration_minutes": 40},
        learning_objectives=[{"objective_id": "obj-1", "description": "解释长乘宽", "evidence_of_achievement": "正确使用平方厘米"}],
        procedure_steps=[{
            "step_id": f"step-{index}", "stage": stage, "duration_minutes": minutes,
            "teacher_actions": [f"教师活动{index}"], "student_actions": [f"学生任务{index}"],
            "assessment": f"评价证据{index}",
        } for index, (stage, minutes) in enumerate(zip(
            ["情境导入", "动手探究", "交流归纳", "巩固练习", "总结评价", "布置作业"],
            [5, 12, 8, 9, 4, 2]), start=1)],
        homework="测量书本封面并说明误差",
        template_extensions={
            "extra_notes": ["备注01｜保留中文与emoji🙂", "备注02\n保留单元格换行", "备注03", "备注04", "备注05", "备注06"],
            "acceptance_end_marker": "LIVE-S1-END",
        },
    )


def text_of(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ])


def test_best_artifact_retains_six_notes_and_core_content(lesson, tmp_path):
    before = lesson.model_dump_json()
    version = LessonPlanVersion(
        version_id="v0", iteration=0, document=lesson,
        created_by_profile_id="offline", created_at=datetime(2026, 10, 15, tzinfo=timezone.utc),
        document_hash=hashlib.sha256(before.encode()).hexdigest(),
        rule_check_report={"report_id": "rules", "plan_id": "plan", "passed": True,
                           "total_duration_minutes": 40, "expected_duration_minutes": 40},
    )
    result = PipelineResult(
        run_id="run", task_id="task", experiment_id="offline", method_id="fixture",
        status="completed", stop_reason="quality_passed", best_version_id="v0",
        last_version_id="v0", versions=[version],
    )
    exported = export_artifacts(result, tmp_path, include_docx=True)
    assert all(record.status == "ok" for record in exported.artifacts.artifacts)
    text = text_of(tmp_path / "best_lesson_plan.docx")
    for expected in ["备注01｜保留中文与emoji🙂", "备注02\n保留单元格换行", "备注03", "备注04", "备注05", "备注06", "LIVE-S1-END"]:
        assert text.count(expected) == 1
    for expected in ["解释长乘宽", "正确使用平方厘米", "测量书本封面并说明误差"]:
        assert expected in text
    process = Document(tmp_path / "best_lesson_plan.docx").tables[2]
    assert len(process.rows) == 7
    for index in range(1, 7):
        assert f"教师活动{index}" in process.rows[index].cells[0].text
        assert f"学生任务{index}" in process.rows[index].cells[1].text
        assert f"评价证据{index}" in process.rows[index].cells[2].text
    assert lesson.model_dump_json() == before
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    docx_record = next(item for item in manifest["artifacts"] if item["format"] == "docx")
    assert docx_record["sha256"] == hashlib.sha256((tmp_path / "best_lesson_plan.docx").read_bytes()).hexdigest()


@pytest.mark.parametrize("extensions, expected", [
    ({"教师说明": "第一行\n第二行🙂"}, ["教师说明", "第一行\n第二行🙂"]),
    ({"材料": {"数量": 0, "已核验": False, "待补": None, "说明": ["中文🙂", "换行\n内容"]}},
     ['"数量": 0', '"已核验": false', '"待补": null', "中文🙂", '换行\\n内容']),
    ({"人数": 0, "已核验": False}, ["人数", "0", "已核验", "false"]),
])
def test_direct_export_preserves_other_json_extension_values(lesson, tmp_path, extensions, expected):
    lesson.template_extensions = extensions
    path = export_docx(lesson, tmp_path / "candidate.docx", run_id="run", best_version_id="v1", last_version_id="v1")
    text = text_of(path)
    for value in expected:
        assert value in text


def test_document_without_extensions_has_no_empty_extension_section(lesson, tmp_path):
    lesson.template_extensions = {}
    path = export_docx(lesson, tmp_path / "plain.docx", run_id="run", best_version_id="v0", last_version_id="v0")
    assert "补充内容" not in text_of(path)
