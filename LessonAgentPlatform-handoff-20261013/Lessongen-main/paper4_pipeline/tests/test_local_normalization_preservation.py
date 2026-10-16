"""The local boundary begins at Word extraction, before any model rewrite."""
from types import SimpleNamespace

import pytest

from paper4_pipeline.agents.protocols import AgentCallMetadata
from paper4_pipeline.domain.models import TokenUsage
from paper4_pipeline.web_api.schemas import EngineLessonInput, NormalizationModelOutput, RawLessonDocument
from test_local_optimization_scope import baseline


def source(minutes=45):
    paragraphs = ["教学目标：解释面积", "作业：原作业，保持逐字不变", "备注：原备注"]
    cells = ["环节", "时间", "教师活动", "学生活动", "评价", "探究", str(minutes), "引导比较", "测量", "解释依据"]
    return RawLessonDocument(source_sha256="a" * 64, original_filename="original.docx",
        paragraphs=[{"locator": f"paragraph:{i}", "text": text} for i, text in enumerate(paragraphs)],
        tables=[{"locator": "table:1", "rows": 2, "columns": 5, "cells": [
            {"locator": f"table:1/cell:{i}", "row": i // 5, "column": i % 5, "text": text}
            for i, text in enumerate(cells)]}])


def task():
    return EngineLessonInput(mode="optimize", subject="数学", grade="五年级", topic="面积",
        optimization_scope={"mode": "local", "editable_paths": ["/procedure_steps"]})


def response(value):
    return SimpleNamespace(value=value, usage=TokenUsage(), estimated_cost=0,
        metadata=AgentCallMetadata(provider="local-fixture", model_name="fixture", prompt_id="fixture",
                                   prompt_version="1", prompt_sha256="0" * 64))


def normalize(document):
    from paper4_pipeline.web_api.normalizer import DocxNormalizer
    value = NormalizationModelOutput(lesson_plan=document,
        metadata_provenance={key: "user" for key in ["subject", "grade", "topic", "duration_minutes"]},
        raw_locator_coverage={"claimed_mapped": ["all"]})
    def invoke(**kwargs):
        kwargs["result_validator"](value)
        return response(value)
    normalizer = object.__new__(DocxNormalizer)
    normalizer.provider = SimpleNamespace(invoke_structured=invoke)
    normalizer.prompt = SimpleNamespace(version="fixture")
    return normalizer.normalize(source(), task(), task_id="task", plan_id="plan")


def test_faithful_local_normalization_records_program_checked_source_coverage():
    result = normalize(baseline())
    assert result.lesson_plan.homework == baseline().homework
    assert "/homework" in result.raw_locator_coverage["paragraph:1"]
    assert "claimed_mapped" not in result.raw_locator_coverage


@pytest.mark.parametrize("mutation", ["delete", "append", "move", "invent", "copy_other_section"])
def test_local_normalizer_cannot_rewrite_or_hide_source_before_baseline(mutation):
    document = baseline()
    if mutation == "delete": document.homework = ""
    if mutation == "append": document.homework += "；新加题目"
    if mutation == "move":
        document.reflection += document.homework
        document.homework = ""
    if mutation == "invent": document.board_design = "新编的板书"
    if mutation == "copy_other_section": document.homework += "；原备注"
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        normalize(document)


def test_local_segmented_import_keeps_original_duration_instead_of_proportional_rewrite():
    from paper4_pipeline.web_api.segmented_normalizer import SegmentedDocxNormalizer, OverviewOutput, ActivityOutput
    document = baseline()
    document.procedure_steps[0].duration_minutes = 40
    overview = OverviewOutput(learning_objectives=document.learning_objectives,
                              homework=document.homework, reflection=document.reflection)
    activities = ActivityOutput(procedure_steps=document.procedure_steps)
    def invoke(**kwargs):
        return response(overview if kwargs["output_schema"] is OverviewOutput else activities)
    provider = SimpleNamespace(invoke_structured=invoke, settings=SimpleNamespace(provider="fixture", model_name="fixture"))
    result = SegmentedDocxNormalizer(provider).normalize(source(40), task(), task_id="task", plan_id="plan")
    assert sum(step.duration_minutes for step in result.lesson_plan.procedure_steps) == 40
    assert any("未调整" in warning for warning in result.warnings)


def test_source_time_four_is_not_satisfied_by_canonical_forty_five():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        checked_source_coverage(source(4), baseline(), task())


def test_duplicate_table_rows_cannot_disappear_during_local_normalization():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    table = raw.tables[0]
    table.cells.extend(cell.model_copy(update={"row": 2, "locator": cell.locator + "-duplicate"})
                       for cell in list(table.cells) if cell.row == 1)
    table.rows = 3
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        checked_source_coverage(raw, baseline(), task())


def test_local_normalizer_cannot_repeat_an_existing_outside_paragraph():
    document = baseline()
    document.reflection += ";" + document.reflection
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        normalize(document)


def test_repeated_source_paragraph_cannot_collapse_to_one_copy():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.paragraphs.append(raw.paragraphs[-1].model_copy(update={"locator": "paragraph:duplicate"}))
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        checked_source_coverage(raw, baseline(), task())


@pytest.mark.parametrize("separate_paragraphs", [False, True])
def test_outside_field_source_order_cannot_change(separate_paragraphs):
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.paragraphs[-1].text = "备注：先比较" if separate_paragraphs else "备注：先比较；后测量"
    if separate_paragraphs:
        raw.paragraphs.append(raw.paragraphs[-1].model_copy(update={"locator": "paragraph:next", "text": "后测量"}))
    document = baseline()
    document.reflection = "后测量；先比较"
    with pytest.raises(ValueError, match="原稿|保全|顺序"):
        checked_source_coverage(raw, document, task())
    document.reflection = "先比较\n后测量"
    assert checked_source_coverage(raw, document, task())


def test_list_items_within_one_field_keep_their_source_order():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.tables[0].cells[7].text = "先比较；后测量"
    document = baseline()
    document.procedure_steps[0].teacher_actions = ["后测量", "先比较"]
    with pytest.raises(ValueError, match="原稿|保全|顺序"):
        checked_source_coverage(raw, document, task())
    document.procedure_steps[0].teacher_actions = ["先比较", "后测量"]
    assert checked_source_coverage(raw, document, task())


def test_combined_explicit_metadata_labels_are_bound_individually():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.paragraphs.insert(0, raw.paragraphs[0].model_copy(update={
        "locator": "paragraph:metadata", "text": "学科：数学；年级：五年级；课时：45分钟。"}))
    assert checked_source_coverage(raw, baseline(), task())
    raw.paragraphs[0].text = "学科：五年级；年级：数学；课时：45分钟。"
    with pytest.raises(ValueError, match="原稿|保全|栏目"):
        checked_source_coverage(raw, baseline(), task())


def test_two_explicit_content_labels_in_one_paragraph_keep_separate_fields():
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.paragraphs.append(raw.paragraphs[-1].model_copy(update={
        "locator": "paragraph:combined", "text": "学情：已会测量。准备：方格纸"}))
    document = baseline()
    document.student_analysis = "已会测量。"
    document.resources = [{"resource_id": "paper", "name": "方格纸"}]
    assert checked_source_coverage(raw, document, task())
    document.student_analysis = "已会测量。方格纸"
    document.resources = []
    with pytest.raises(ValueError, match="原稿|保全|栏目"):
        checked_source_coverage(raw, document, task())


def test_extra_table_notes_are_preserved_with_their_source_row_and_visible_in_word(tmp_path):
    from docx import Document
    from paper4_pipeline.exporters.docx_exporter import export_docx
    from paper4_pipeline.web_api.schemas import RawTableCell
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    table = raw.tables[0]
    table.columns = 6
    table.cells.extend([
        RawTableCell(locator="notes:header", row=0, column=5, text="额外备注"),
        RawTableCell(locator="notes:row1", row=1, column=5, text="备注01\n保留中文与emoji🙂"),
    ])
    document = baseline()
    checked = checked_source_coverage(raw, document, task())
    extra = document.template_extensions["source_table_columns"][0]
    assert extra == {"table_locator": "table:1", "column_index": 5, "heading": "额外备注",
                     "rows": [{"row_index": 1, "stage": "探究", "text": "备注01\n保留中文与emoji🙂"}]}
    assert checked["notes:row1"] == ["/template_extensions/source_table_columns/0/rows/0/text"]
    path = export_docx(document, tmp_path / "notes.docx", run_id="r", best_version_id="v0", last_version_id="v0")
    exported = Document(path)
    table = next(table for table in exported.tables if "额外备注" in table.cell(0, 1).text)
    assert table.cell(1, 0).text == "探究"
    assert table.cell(1, 1).text == "备注01\n保留中文与emoji🙂"


def test_model_cannot_move_extra_table_note_into_another_field():
    from paper4_pipeline.web_api.schemas import RawTableCell
    from paper4_pipeline.web_api.source_preservation import checked_source_coverage
    raw = source()
    raw.tables[0].columns = 6
    raw.tables[0].cells.extend([
        RawTableCell(locator="notes:header", row=0, column=5, text="额外备注"),
        RawTableCell(locator="notes:row1", row=1, column=5, text="独立备注"),
    ])
    document = baseline()
    document.reflection += ";独立备注"
    with pytest.raises(ValueError, match="原稿|保全|来源"):
        checked_source_coverage(raw, document, task())
