"""Offline checks of real export recovery and upload rejection boundaries."""
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import json

import pytest
from docx import Document

from paper4_pipeline.domain.models import LessonPlanDocument, LessonPlanVersion, PipelineResult
from paper4_pipeline.exporters.manifest import export_artifacts
from paper4_pipeline.web_api.docx_ingestion import DocxIngestionError, extract_docx, inspect_docx
from paper4_pipeline.web_api.generated_docx import restore_generated_docx
from paper4_pipeline.web_api.schemas import EngineLessonInput
from paper4_pipeline.web_api.settings import EngineSettings


@pytest.fixture
def native_export(tmp_path):
    config = EngineSettings(state_root=tmp_path / 'state', artifacts_root=tmp_path / 'artifacts',
                            config_path=tmp_path / 'unused-config.json')
    document = LessonPlanDocument(plan_id='original-plan', task_id='original-task',
        metadata={'subject': '数学', 'grade': '五年级', 'topic': '面积', 'duration_minutes': 45},
        procedure_steps=[{'step_id': 'step-1', 'stage': '探究', 'duration_minutes': 45,
                          'teacher_actions': ['引导比较'], 'student_actions': ['测量并解释'],
                          'assessment': '观察解释是否完整'}], homework='测量三角形')
    version = LessonPlanVersion(version_id='version-1', iteration=0, document=document,
        created_by_profile_id='offline-fixture', created_at=datetime(2026, 10, 14, tzinfo=timezone.utc),
        document_hash=hashlib.sha256(document.model_dump_json().encode()).hexdigest(),
        rule_check_report={'report_id': 'report-1', 'plan_id': 'original-plan', 'passed': True,
                           'total_duration_minutes': 45, 'expected_duration_minutes': 45})
    result = PipelineResult(run_id='run-1', task_id='original-task', experiment_id='offline',
        method_id='fixture', status='completed', stop_reason='quality_passed',
        best_version_id='version-1', last_version_id='version-1', versions=[version])
    folder = config.artifacts_root / 'run-1'
    exported = export_artifacts(result, folder, include_docx=True)
    assert all(record.status == 'ok' for record in exported.artifacts.artifacts)
    path = folder / 'best_lesson_plan.docx'
    raw = extract_docx(path, original_filename=path.name, settings=config)
    task = EngineLessonInput(mode='optimize', subject='数学', grade='五年级', topic='面积')
    return config, folder, raw, task


def restore(native_export):
    config, _, raw, task = native_export
    return restore_generated_docx(raw, task, task_id='new-task', plan_id='new-plan',
                                  artifacts_root=config.artifacts_root)


def test_matching_real_export_recovers_canonical_content_without_model(native_export):
    restored = restore(native_export)
    assert restored is not None
    assert restored.model_metadata['provider'] == 'verified_export'
    assert restored.model_metadata['attempts'] == 0
    assert restored.model_metadata['source_run_id'] == 'run-1'
    assert restored.lesson_plan.task_id == 'new-task'
    assert restored.lesson_plan.plan_id == 'new-plan'
    assert restored.lesson_plan.procedure_steps[0].teacher_actions == ['引导比较']
    assert restored.lesson_plan.procedure_steps[0].student_actions == ['测量并解释']
    assert restored.lesson_plan.homework == '测量三角形'


@pytest.mark.parametrize('artifact', ['best_lesson_plan.docx', 'best_lesson_plan.json'])
def test_changed_artifact_bytes_cannot_reuse_original_manifest(native_export, artifact):
    _, folder, _, _ = native_export
    assert restore(native_export) is not None
    path = folder / artifact
    if path.suffix == '.json':
        data = json.loads(path.read_text(encoding='utf-8'))
        data['lesson_plan']['homework'] = '未经确认的作业'
        path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    else:
        document = Document(path)
        document.add_paragraph('未经确认的补充内容')
        document.save(path)
    assert restore(native_export) is None


def test_existing_export_files_without_manifest_are_untrusted(native_export):
    _, folder, _, _ = native_export
    assert restore(native_export) is not None
    (folder / 'manifest.json').unlink()
    assert restore(native_export) is None


def test_non_zip_docx_is_rejected_before_parsing(tmp_path):
    path = tmp_path / 'invalid.docx'
    path.write_bytes(b'this is not an OOXML document')
    config = EngineSettings(tmp_path / 'state', tmp_path / 'artifacts', tmp_path / 'unused.json')
    with pytest.raises(DocxIngestionError) as failure:
        extract_docx(path, original_filename=path.name, settings=config)
    assert failure.value.code == 'INVALID_DOCX'


def test_upload_limit_accepts_exact_size_and_rejects_one_byte_over(native_export):
    config, folder, _, _ = native_export
    path = folder / 'best_lesson_plan.docx'
    size = path.stat().st_size
    inspect_docx(path, replace(config, max_upload_bytes=size))
    with pytest.raises(DocxIngestionError) as failure:
        inspect_docx(path, replace(config, max_upload_bytes=size - 1))
    assert failure.value.code == 'DOCX_TOO_LARGE'
