"""Use the real SDK with controlled HTTP responses; never call an external model."""
import json
from importlib import import_module

import pytest
import openai

import providers

# SDK 2 uses httpx; SDK 3 uses httpx2. Mock the SDK's actual client transport,
# including environments where both libraries are installed, so no model request
# can accidentally bypass the controlled response handler.
httpx = import_module(openai.DefaultHttpxClient.__mro__[1].__module__.split('.')[0])


def completion(usage=None):
    data = {'id': 'completion-test', 'object': 'chat.completion', 'created': 1,
            'model': 'test-model', 'choices': [{'index': 0, 'finish_reason': 'stop',
            'message': {'role': 'assistant', 'content': 'controlled answer'}}]}
    if usage is not None:
        data['usage'] = usage
    return data


def provider(tmp_path, monkeypatch, handler):
    monkeypatch.setattr(providers, 'load_local_key', lambda name: '')
    monkeypatch.setattr(httpx.HTTPTransport, 'handle_request', handler)
    monkeypatch.setattr(openai._base_client.time, 'sleep', lambda seconds: None)
    return providers.create_provider('deepseek', {
        'api_keys': {'deepseek': 'test-secret'}, 'llm': {'model': 'test-model'},
        'output': {'dir': str(tmp_path)}, '_run_id': 'run-123'})


def records(tmp_path):
    path = tmp_path / 'model_usage.jsonl'
    assert path.exists(), 'Every model attempt must have a durable usage ledger'
    return [json.loads(line) for line in path.read_text(encoding='utf-8').splitlines()]


def test_attempt_artifacts_keep_the_canonical_usage_ledger(tmp_path,monkeypatch):
    monkeypatch.setattr(providers, 'load_local_key', lambda name: '')
    monkeypatch.setattr(httpx.HTTPTransport, 'handle_request',
        lambda self, request: httpx.Response(200,json=completion(),request=request))
    canonical = tmp_path / 'output/model_usage.jsonl'
    client = providers.create_provider('deepseek', {
        'api_keys': {'deepseek':'test-secret'}, '_run_id':'run-contained',
        'output': {'dir': str(tmp_path/'attempts/unique/output')},
        '_usage_ledger_path': str(canonical)})
    assert client.chat_json('private prompt','private content') == 'controlled answer'
    rows = [json.loads(line) for line in canonical.read_text().splitlines()]
    assert {row['task_id'] for row in rows} == {'run-contained'}
    assert any(row['event']=='request_finished' for row in rows)
    assert not (tmp_path/'attempts/unique/output/model_usage.jsonl').exists()


def test_sdk_retry_records_every_attempt_and_preserves_unknown_usage(tmp_path, monkeypatch):
    attempts = []
    def transport(self, request):
        attempts.append(request)
        if len(attempts) == 1:
            return httpx.Response(500, json={'error': {'message': 'private error'}},
                                  headers={'x-request-id': 'provider-failed'}, request=request)
        return httpx.Response(200, json=completion({'prompt_tokens': 11, 'completion_tokens': 7,
            'total_tokens': 18, 'prompt_tokens_details': {'cached_tokens': 4},
            'secret': 'must not enter ledger'}), headers={'x-request-id': 'provider-ok'}, request=request)
    client = provider(tmp_path, monkeypatch, transport)
    assert client.chat_json('private-system-prompt', 'private-lesson-content') == 'controlled answer'
    rows = records(tmp_path)
    finished = [row for row in rows if row['event'] == 'request_finished']
    assert len(finished) == 2
    assert [row['sdk_retry_index'] for row in finished] == [0, 1]
    assert [row['status'] for row in finished] == ['HTTP_ERROR', 'SUCCEEDED']
    assert finished[0]['usage'] is None
    assert finished[1]['usage'] == {'prompt_tokens': 11, 'completion_tokens': 7,
        'total_tokens': 18, 'prompt_tokens_details': {'cached_tokens': 4}}
    assert {row['task_id'] for row in rows} == {'run-123'}
    assert len({row['call_id'] for row in rows}) == 1
    assert len({row['request_id'] for row in finished}) == 2
    assert finished[1]['provider_request_id'] == 'provider-ok'
    text = json.dumps(rows)
    for secret in ['test-secret', 'private-system-prompt', 'private-lesson-content', 'private error', 'must not enter ledger']:
        assert secret not in text


def test_exhausted_sdk_timeouts_are_not_counted_as_zero_tokens(tmp_path, monkeypatch):
    def transport(self, request):
        raise httpx.ReadTimeout('private request body', request=request)
    client = provider(tmp_path, monkeypatch, transport)
    with pytest.raises(openai.APITimeoutError):
        client.chat('system', 'input')
    finished = [row for row in records(tmp_path) if row['event'] == 'request_finished']
    assert len(finished) == 3
    assert [row['sdk_retry_index'] for row in finished] == [0, 1, 2]
    assert all(row['status'] == 'TRANSPORT_ERROR' and row['usage'] is None for row in finished)
    assert all(row['error_type'] == 'ReadTimeout' for row in finished)


def test_success_without_usage_remains_unknown(tmp_path, monkeypatch):
    client = provider(tmp_path, monkeypatch,
                      lambda self, request: httpx.Response(200, json=completion(), request=request))
    assert client.chat('system', 'input') == 'controlled answer'
    finished = [row for row in records(tmp_path) if row['event'] == 'request_finished']
    assert len(finished) == 1
    assert finished[0]['status'] == 'SUCCEEDED' and finished[0]['usage'] is None


def test_application_empty_output_retries_share_operation_and_run(tmp_path, monkeypatch):
    from docx import Document
    import platform_api
    from service import run_annotation, AnnotationError
    body = completion({'prompt_tokens': 2, 'completion_tokens': 0, 'total_tokens': 2})
    body['choices'][0]['message']['content'] = ''
    provider(tmp_path, monkeypatch, lambda self, request: httpx.Response(200, json=body, request=request))
    input_dir = tmp_path / 'input'
    input_dir.mkdir()
    doc = Document()
    doc.add_paragraph('数学 三年级 长方形面积教学方案：使用单位正方形覆盖长方形。')
    doc.save(input_dir / 'lesson.docx')
    config = {'api_keys': {'deepseek': 'test-secret'}, 'llm': {'provider': 'deepseek', 'model': 'test-model'},
              'data': {'plans_dir': str(input_dir)}, 'prompt': {'max_examples': 0},
              'output': {'dir': str(tmp_path)}, '_run_id': 'run-application'}
    with pytest.raises(AnnotationError, match='API'):
        run_annotation(str(platform_api.ROOT), config, 'lesson.docx', criteria_type='student')
    finished = [row for row in records(tmp_path) if row['event'] == 'request_finished']
    assert len(finished) == 3
    assert {row['operation'] for row in finished} == {'annotation'}
    assert {row['task_id'] for row in finished} == {'run-application'}
    assert len({row['operation_id'] for row in finished}) == 1
    assert [row['application_attempt'] for row in finished] == [1, 2, 3]
    assert finished[0]['retry_of_call_id'] is None
    assert finished[1]['retry_of_call_id'] == finished[0]['call_id']
    assert finished[2]['retry_of_call_id'] == finished[1]['call_id']


def test_platform_run_binds_its_run_id_for_provider_metering(tmp_path, monkeypatch):
    import platform_api
    monkeypatch.setattr(platform_api, 'RUN_ROOT', tmp_path)
    request = platform_api.AnnotationRequest(lesson_id='l', version_id='v', subject='数学',
                                            grade='三年级', topic='面积', content='内容')
    run_id, _, _ = platform_api._initialize_run(request)
    captured = []
    def evaluate(base_dir, config, *args, **kwargs):
        captured.append(config.get('_run_id'))
        raise RuntimeError('controlled stop before any model call')
    monkeypatch.setattr(platform_api, 'execute_annotation_worker', evaluate)
    platform_api._execute(run_id, request)
    assert captured == [run_id]


def test_suggestion_repairs_have_their_own_operation(tmp_path, monkeypatch):
    import suggester
    client = provider(tmp_path, monkeypatch,
                      lambda self, request: httpx.Response(200, json=completion(), request=request))
    monkeypatch.setattr(suggester, 'load_knowledge_base', lambda: ['controlled source'])
    monkeypatch.setattr(suggester, '_match_sources', lambda *args: {
        '来源类型': 'controlled', '教学法参考': '', '理论条目': 'controlled source', '课标原文': ''})
    # Existing suggestion diagnostics must also stay in this test's temporary tree.
    monkeypatch.setattr(suggester, '__file__', str(tmp_path / 'suggester.py'))
    result = {'维度评分': [{'维度名称': '目标', '得分': 1, '满分': 3,
                          '批注列表': [{'评价': 'C', '位置': '目标', '具体分析': '缺少可检查目标'}]}]}
    suggester.generate_suggestions(result, 'controlled lesson', client)
    finished = [row for row in records(tmp_path) if row['event'] == 'request_finished']
    assert len(finished) == 2
    assert {row['operation'] for row in finished} == {'suggestions'}
    assert [row['application_attempt'] for row in finished] == [1, 2]
    assert finished[1]['retry_of_call_id'] == finished[0]['call_id']


def test_accounting_disk_failure_does_not_retry_successful_model_request(tmp_path, monkeypatch, caplog):
    from pathlib import Path
    calls = []
    def transport(self, request):
        calls.append(request)
        return httpx.Response(200, json=completion(), request=request)
    client = provider(tmp_path, monkeypatch, transport)
    original_open = Path.open
    def blocked(path, *args, **kwargs):
        if path.name == 'model_usage.jsonl':
            raise OSError('private directory name')
        return original_open(path, *args, **kwargs)
    monkeypatch.setattr(Path, 'open', blocked)
    assert client.chat('system', 'input') == 'controlled answer'
    assert len(calls) == 1
    assert 'MODEL_USAGE_WRITE_FAILED' in caplog.text and 'private directory name' not in caplog.text


def test_successful_annotation_and_suggestion_repairs_are_separate_operations(tmp_path, monkeypatch):
    from docx import Document
    import platform_api
    import suggester
    from service import run_annotation
    score = {'维度评分': [{'维度名称': '教学目标', '得分': 1, '满分': 10,
                          '批注列表': [{'评价': 'C', '位置': '目标', '具体分析': '缺少可检查目标'}]}],
             '评分汇总': {'总分': 1}, '总体评价': 'controlled review'}
    contents = iter([json.dumps(score, ensure_ascii=False), 'invalid suggestions',
                     '{"suggestions": [{"id": 1, "具体建议": "add an exit question"}]}'])
    def transport(self, request):
        body = completion({'prompt_tokens': 3, 'completion_tokens': 2, 'total_tokens': 5})
        body['choices'][0]['message']['content'] = next(contents)
        return httpx.Response(200, json=body, request=request)
    provider(tmp_path, monkeypatch, transport)
    monkeypatch.setattr(suggester, 'load_knowledge_base', lambda: ['controlled source'])
    monkeypatch.setattr(suggester, '_match_sources', lambda *args: {
        '来源类型': 'controlled', '教学法参考': '', '理论条目': 'controlled source', '课标原文': ''})
    monkeypatch.setattr(suggester, '__file__', str(tmp_path / 'suggester.py'))
    input_dir = tmp_path / 'input'
    input_dir.mkdir()
    doc = Document()
    doc.add_paragraph('数学 三年级 长方形面积：使用单位正方形覆盖长方形。')
    doc.save(input_dir / 'lesson.docx')
    config = {'api_keys': {'deepseek': 'test-secret'}, 'llm': {'provider': 'deepseek', 'model': 'test-model'},
              'data': {'plans_dir': str(input_dir)}, 'prompt': {'max_examples': 0},
              'output': {'dir': str(tmp_path)}, '_run_id': 'run-success-and-repair'}
    outcome = run_annotation(str(platform_api.ROOT), config, 'lesson.docx', criteria_type='student', suggest=True)
    assert outcome['result']['维度评分'][0]['批注列表'][0]['细化建议'] == 'add an exit question'
    finished = [row for row in records(tmp_path) if row['event'] == 'request_finished']
    assert len(finished) == 3
    review, initial_suggestion, repaired_suggestion = finished
    assert [row['operation'] for row in finished] == ['annotation', 'suggestions', 'suggestions']
    assert [row['application_attempt'] for row in finished] == [1, 1, 2]
    assert initial_suggestion['operation_id'] != review['operation_id']
    assert initial_suggestion['parent_operation_id'] == review['operation_id']
    assert initial_suggestion['retry_of_call_id'] is None
    assert repaired_suggestion['retry_of_call_id'] == initial_suggestion['call_id']
