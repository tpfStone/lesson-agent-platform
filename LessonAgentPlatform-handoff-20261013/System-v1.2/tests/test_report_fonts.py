"""New F1 reports use explicit editable fonts; no source document is rewritten."""
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn

from annotator import generate_docx


def fonts_of(run):
    r_pr = run._element.rPr
    assert r_pr is not None
    fonts = r_pr.find(qn("w:rFonts"))
    assert fonts is not None
    return {fonts.get(qn("w:" + slot)) for slot in ("ascii", "hAnsi", "eastAsia", "cs")}


def test_generated_report_fonts_cover_titles_body_tables_and_symbols(tmp_path):
    result = {"教案基本信息": {"课题名称": "面积 Area 12", "学科": "数学", "年级": "五年级"},
              "评分汇总": {"总分": 3}, "总体评价": "原文引用保留，12 cm²。",
              "维度评分": [{"维度名称": "测试维度", "满分": 5, "得分": 3, "等级": "B",
                            "得分理由": "说明依据", "批注列表": [{"评价": "B", "位置": "表1",
                            "原文引用": "原稿🙂", "具体分析": "补充观察", "建议": "说明比较方法",
                            "依据来源": "来源1", "教学法参考": "操作观察", "细化建议": "先测量"}]}]}
    before = deepcopy(result)
    path = tmp_path / "report.docx"
    generate_docx(result, str(path))
    assert result == before
    doc = Document(path)
    assert not doc.paragraphs[-1]._element.xpath('.//w:br[@w:type="page"]')
    def check(r_pr, expected):
        assert r_pr is not None
        fonts = r_pr.find(qn("w:rFonts"))
        assert fonts is not None
        assert {fonts.get(qn("w:" + slot)) for slot in ("ascii", "hAnsi", "eastAsia", "cs")} == {expected}
        assert not any(key.lower().endswith("theme") for key in fonts.attrib)
    for paragraph in doc.paragraphs:
        expected = "SimHei" if paragraph.style.name == "Title" or paragraph.style.name.startswith("Heading") else "SimSun"
        for run in paragraph.runs:
            if run.text:
                check(run._element.rPr, "Segoe UI Emoji" if run.text.strip() in {"📚", "🏫", "💡"} else expected)
    for table in doc.tables:
        assert all(run.bold for cell in table.rows[0].cells for paragraph in cell.paragraphs for run in paragraph.runs if run.text)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text:
                            check(run._element.rPr, "SimSun")
    for name, family in [("Normal", "SimSun"), ("Title", "SimHei"), ("Heading 1", "SimHei"), ("Heading 2", "SimHei")]:
        check(doc.styles[name].element.rPr, family)
    check(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")), "SimSun")


def test_batch_report_uses_the_same_heading_body_and_table_fonts(tmp_path):
    from batch_reporting import write_batch_quality_report
    paths = write_batch_quality_report([{"result": {"教案基本信息": {"课题名称": "面积"},
        "评分汇总": {"总分": 80}, "维度评分": []}}], tmp_path)
    doc = Document(paths["quality_docx"])
    for paragraph in doc.paragraphs:
        expected = "SimHei" if paragraph.style.name == "Title" or paragraph.style.name.startswith("Heading") else "SimSun"
        for run in paragraph.runs:
            if run.text:
                assert fonts_of(run) == {expected}
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text:
                            assert fonts_of(run) == {"SimSun"}


def test_native_comments_and_unlocated_appendix_style_only_new_text(tmp_path):
    import hashlib
    from docx_annotator import generate_annotated_docx
    source = tmp_path / "original.docx"
    doc = Document()
    doc.styles["Normal"].font.name = "Courier New"
    original = doc.add_paragraph().add_run("保留原文引用位置")
    original.font.name = "KaiTi"
    original._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "KaiTi")
    original_properties = original._element.rPr.xml
    doc.save(source)
    original_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    result = {"维度评分": [{"维度名称": "测试", "批注列表": [
        {"评价": "B", "原文引用": "保留原文引用位置", "具体分析": "这里新增的分析"},
        {"评价": "C", "原文引用": "ZZZ无法匹配", "具体分析": "未定位的新分析"}]}]}
    output = tmp_path / "annotated.docx"
    assert generate_annotated_docx(str(source), result, str(output)) == str(output)
    annotated = Document(output)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_sha
    assert annotated.styles["Normal"].font.name == "Courier New"
    assert annotated.paragraphs[0].runs[0]._element.rPr.xml == original_properties
    comments = list(annotated.comments)
    assert len(comments) == 2
    for comment in comments:
        for paragraph in comment.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    assert fonts_of(run) == {"SimSun"}
    for paragraph in annotated.paragraphs[1:]:
        for run in paragraph.runs:
            if run.text:
                assert fonts_of(run) == {"SimHei" if run.text == "未定位批注汇总" else "SimSun"}
